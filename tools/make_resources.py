#!/usr/bin/env python3
"""Builds every free resource into downloads/, recalculates Excel formulas, renders preview
images into images/resources/ and writes downloads/resources.json for the Resources page.
Run when resources change:  python3 tools/make_resources.py"""
import json, os, subprocess, sys, html
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import resource_pdfs, resource_xlsx
from PIL import Image

DL = os.path.join(ROOT, 'downloads'); PV = os.path.join(ROOT, 'images', 'resources')
os.makedirs(DL, exist_ok=True); os.makedirs(PV, exist_ok=True)
RECALC = '/mnt/skills/public/xlsx/scripts/recalc.py'
TMP = os.path.join(ROOT, 'tools', '_tmp'); os.makedirs(TMP, exist_ok=True)


def size(p):
    b = os.path.getsize(p)
    return '%d KB' % max(1, round(b / 1024)) if b < 1024 * 1024 else '%.1f MB' % (b / 1048576)


def pdf_preview(pdf, slug):
    base = os.path.join(TMP, slug)
    subprocess.run(['pdftoppm', '-r', '110', '-png', '-f', '1', '-l', '1', pdf, base], check=True)
    png = [f for f in os.listdir(TMP) if f.startswith(slug + '-') and f.endswith('.png')][0]
    im = Image.open(os.path.join(TMP, png)).convert('RGB')
    w, h = im.size
    im = im.crop((0, 0, w, int(w * 0.78)))
    im.thumbnail((720, 720))
    im.save(os.path.join(PV, slug + '.webp'), quality=82, method=6)
    os.remove(os.path.join(TMP, png))


def xlsx_preview(xlsx, slug, meta):
    """Render the main sheet as a styled HTML table, then screenshot it."""
    import openpyxl
    from playwright.sync_api import sync_playwright
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb['Status'] if meta['file'] == 'reporting-timeliness-tracker.xlsx' else (wb[wb.sheetnames[1]] if meta['file'] != 'sample-epi-xlsform.xlsx' else wb['survey'])
    hdr_row = 4 if meta['file'] != 'sample-epi-xlsform.xlsx' else 1
    pick = {'stock-tracker.xlsx': 'ABCDEGIJ', 'indicator-tracker.xlsx': 'ADEFIJL', 'data-cleaning-log.xlsx': 'ACDEFGI',
            'it-asset-register.xlsx': 'ABEFHIK', 'sample-epi-xlsform.xlsx': 'ABCEFG',
            'reporting-timeliness-tracker.xlsx': 'ABCDEFGHI', 'training-tracker.xlsx': 'AEFGHKLMNO'}[meta['file']]
    from openpyxl.utils import column_index_from_string as ci
    colsel = [ci(x) for x in pick]
    heads = [ws.cell(row=hdr_row, column=c).value or '' for c in colsel]
    rows = []
    for r in range(hdr_row + 1, hdr_row + 9):
        vals = []
        for c in colsel:
            v = ws.cell(row=r, column=c).value
            if hasattr(v, 'strftime'):
                v = v.strftime('%d %b %Y')
            elif isinstance(v, float):
                v = ('%.0f%%' % (v * 100)) if 0 < v < 1.5 and any(w in str(heads[colsel.index(c)]).lower() for w in ('target', 'test')) else ('%.1f' % v).rstrip('0').rstrip('.')
            vals.append('' if v is None else str(v))
        rows.append(vals)
    status = {'REORDER': 'bad', 'Behind': 'bad', 'Expired': 'bad', 'Open': 'bad', 'OVERSTOCK': 'warn', 'Watch': 'warn',
              'Ending soon': 'warn', 'Query sent': 'warn', 'OK': 'ok', 'On track': 'ok', 'Active': 'ok', 'Fixed': 'ok',
              'Missing': 'bad', 'Late': 'warn', 'On time': 'ok', 'Yes': 'ok', 'No': 'bad'}
    def td(v, i, r):
        cls = status.get(v)
        inner = '<span class="chip %s">%s</span>' % (cls, html.escape(v)) if cls else html.escape(v)
        return '<td class="%s">%s</td>' % ('ex' if r < 3 and meta['file'] != 'sample-epi-xlsform.xlsx' else '', inner)
    body = ''.join('<tr><th class="rn">%d</th>%s</tr>' % (i + hdr_row + 1, ''.join(td(v, j, i) for j, v in enumerate(r))) for i, r in enumerate(rows))
    tabs = ''.join('<span class="%s">%s</span>' % ('on' if n == ws.title else '', html.escape(n)) for n in wb.sheetnames)
    page = f'''<html><head><style>
@font-face{{font-family:DM;src:url(file://{ROOT}/fonts/dm-sans-latin-var.woff2)}}
@font-face{{font-family:BG;src:url(file://{ROOT}/fonts/bricolage-grotesque-latin-var.woff2)}}
body{{margin:0;background:#fff;font:13px DM,Arial,sans-serif;color:#0F2430;width:720px;height:562px;overflow:hidden}}
.bar{{height:38px;background:#0B7A5C;color:#fff;display:flex;align-items:center;gap:10px;padding:0 16px;font-weight:700}}
.bar i{{width:18px;height:18px;border-radius:4px;background:#fff;display:inline-block;position:relative}}
.bar i::after{{content:"";position:absolute;inset:4px;background:#0B7A5C;border-radius:1px}}
.bar small{{margin-left:auto;font:500 11px monospace;opacity:.8}}
.fx{{height:30px;border-bottom:1px solid #D3DED9;display:flex;align-items:center;gap:10px;padding:0 12px;color:#475D68;font:12px monospace}}
.fx b{{color:#0B7A5C}}
h1{{font:800 22px BG,Arial;margin:14px 16px 2px}} p{{margin:0 16px 12px;color:#475D68;font-size:12px}}
table{{border-collapse:collapse;margin:0 0 0 0;width:100%;table-layout:fixed}}
th,td{{border:1px solid #D3DED9;padding:7px 8px;text-align:left;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;font-size:12px}}
thead th{{background:#0B7A5C;color:#fff;font-weight:700;white-space:normal;line-height:1.2;height:34px}}
th.rn{{width:26px;background:#F3F6F4;color:#475D68;font-weight:400;text-align:center}}
td.ex{{background:#FBEBC8}}
.chip{{display:inline-block;padding:2px 8px;border-radius:99px;font-weight:700;font-size:11px}}
.bad{{background:#FDE2E1;color:#B42318}}.warn{{background:#FBEBC8;color:#7A4A00}}.ok{{background:#DDF0E8;color:#0B7A5C}}
.tabs{{position:absolute;bottom:0;left:0;right:0;height:30px;background:#F3F6F4;border-top:1px solid #D3DED9;display:flex;gap:2px;padding-left:12px}}
.tabs span{{padding:6px 14px;font-size:11px;color:#475D68}}.tabs .on{{background:#fff;color:#0B7A5C;font-weight:700;border-bottom:2px solid #0B7A5C}}
</style></head><body>
<div class="bar"><i></i>{html.escape(meta['file'])}<small>alhajiimehemoh.github.io</small></div>
<div class="fx"><b>fx</b>{html.escape('Real XLSForm: survey, choices and settings sheets' if meta['file'] == 'sample-epi-xlsform.xlsx' else 'Formulas fill the grey columns automatically')}</div>
<h1>{html.escape(meta['title'])}</h1><p>{html.escape(meta['blurb'])}</p>
<table><thead><tr><th class="rn"></th>{''.join('<th>%s</th>' % html.escape(str(h)) for h in heads)}</tr></thead><tbody>{body}</tbody></table>
<div class="tabs">{tabs}</div></body></html>'''
    hp = os.path.join(TMP, slug + '.html')
    open(hp, 'w').write(page)
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={'width': 720, 'height': 562})
        pg.goto('file://' + hp); pg.wait_for_timeout(300)
        out = os.path.join(TMP, slug + '.png'); pg.screenshot(path=out); b.close()
    Image.open(out).convert('RGB').save(os.path.join(PV, slug + '.webp'), quality=84, method=6)


def main():
    items = []
    for m in resource_pdfs.build_pdfs(DL):
        p = os.path.join(DL, m['file']); slug = m['file'][:-4]
        pages = int([l for l in subprocess.run(['pdfinfo', p], capture_output=True, text=True).stdout.splitlines() if l.startswith('Pages')][0].split()[-1])
        pdf_preview(p, slug)
        items.append({'file': m['file'], 'slug': slug, 'type': 'PDF', 'category': m['category'], 'title': m['title'].replace('&amp;', '&'),
                      'blurb': m['blurb'], 'tags': m['tags'], 'size': size(p), 'pages': pages})
    for m in resource_xlsx.build_xlsx(DL):
        p = os.path.join(DL, m['file']); slug = m['file'][:-5]
        r = subprocess.run([sys.executable, RECALC, p, '90'], capture_output=True, text=True)
        res = json.loads(r.stdout[r.stdout.index('{'):])
        if res.get('status') != 'success':
            raise SystemExit('Formula errors in %s: %s' % (m['file'], res))
        xlsx_preview(p, slug, m)
        items.append({'file': m['file'], 'slug': slug, 'type': 'XLSForm' if 'xlsform' in slug else 'Excel', 'category': m['category'],
                      'title': m['title'], 'blurb': m['blurb'], 'tags': m['tags'], 'size': size(p),
                      'sheets': len(__import__('openpyxl').load_workbook(p, read_only=True).sheetnames)})
    import zipfile
    zp = os.path.join(DL, 'free-resource-kit.zip')
    with zipfile.ZipFile(zp, 'w', zipfile.ZIP_DEFLATED) as z:
        for it in items:
            z.write(os.path.join(DL, it['file']), 'Alhaji-Mehemoh-resource-kit/' + it['file'])
        z.writestr('Alhaji-Mehemoh-resource-kit/README.txt',
                   'Free resource kit by Alhaji Ibrahim Mehemoh\r\n\r\n' +
                   '\r\n'.join('- %s (%s): %s' % (it['title'], it['type'], it['blurb']) for it in items) +
                   '\r\n\r\nFree to use and adapt. Please keep the credit line.\r\nalhajiimehemoh.github.io  |  alhajiimehemoh@gmail.com\r\n')
    json.dump({'items': items, 'kit': {'file': 'free-resource-kit.zip', 'size': size(zp)}}, open(os.path.join(DL, 'resources.json'), 'w'), indent=1)
    for f in os.listdir(TMP):
        os.remove(os.path.join(TMP, f))
    os.rmdir(TMP)
    print('built', len(items), 'resources')


if __name__ == '__main__':
    main()
