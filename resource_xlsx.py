"""Branded Excel resources. Shared template: a 'How to use' cover sheet, a styled main table with
input cells marked, formulas guarded against blanks, validation lists and conditional formatting.
Example rows are labelled as examples."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

GREEN, INK, SOFT, AMBER, AMBER_SOFT, PAPER, RULE, MUTED = '0B7A5C', '0F2430', 'DDF0E8', 'D68A00', 'FBEBC8', 'F3F6F4', 'D3DED9', '475D68'
F = 'Arial'
thin = Side(style='thin', color=RULE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT_FILL = PatternFill('solid', fgColor='FFFFFF')
CALC_FILL = PatternFill('solid', fgColor=PAPER)
EX_FILL = PatternFill('solid', fgColor=AMBER_SOFT)


def fill(c):
    return PatternFill('solid', fgColor=c)


def cover(wb, title, purpose, steps, legend=True):
    ws = wb.active
    ws.title = 'How to use'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 80
    for r in range(1, 4):
        for col in 'ABCD':
            ws[f'{col}{r}'].fill = fill(GREEN)
    ws['B2'] = 'Alhaji Ibrahim Mehemoh  ·  Free resource'
    ws['B2'].font = Font(name=F, size=10, bold=True, color='FFFFFF')
    ws.row_dimensions[2].height = 22
    ws['B5'] = title
    ws['B5'].font = Font(name=F, size=20, bold=True, color=INK)
    ws.row_dimensions[5].height = 30
    ws['B6'] = purpose
    ws['B6'].font = Font(name=F, size=11, color=MUTED)
    ws.merge_cells('B6:C6')
    ws['B6'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[6].height = 34
    r = 8
    ws[f'B{r}'] = 'HOW TO USE'
    ws[f'B{r}'].font = Font(name=F, size=9, bold=True, color=GREEN)
    r += 1
    for i, s in enumerate(steps, 1):
        ws[f'B{r}'] = f'Step {i}'
        ws[f'B{r}'].font = Font(name=F, size=10, bold=True, color=INK)
        ws[f'C{r}'] = s
        ws[f'C{r}'].font = Font(name=F, size=10, color=INK)
        ws[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')
        ws[f'B{r}'].alignment = Alignment(vertical='top')
        ws.row_dimensions[r].height = 30
        r += 1
    if legend:
        r += 1
        ws[f'B{r}'] = 'COLOUR KEY'
        ws[f'B{r}'].font = Font(name=F, size=9, bold=True, color=GREEN)
        r += 1
        for f_, label in [(INPUT_FILL, 'White cells: type your data here'), (CALC_FILL, 'Grey cells: formulas, do not type over them'),
                          (EX_FILL, 'Amber rows: an example. Delete or overwrite it when you start')]:
            ws[f'B{r}'].fill = f_
            ws[f'B{r}'].border = BORDER
            ws[f'C{r}'] = label
            ws[f'C{r}'].font = Font(name=F, size=10, color=INK)
            r += 1
    r += 1
    ws[f'B{r}'] = 'Privacy: do not store patient names, phone numbers or ID numbers in this file.'
    ws[f'B{r}'].font = Font(name=F, size=10, bold=True, color=AMBER)
    ws.merge_cells(f'B{r}:C{r}')
    r += 2
    ws[f'B{r}'] = 'Free to use and adapt. Please keep this credit line.  ·  alhajiimehemoh.github.io  ·  Need it tailored? alhajiimehemoh@gmail.com'
    ws[f'B{r}'].font = Font(name=F, size=9, color=MUTED)
    ws.merge_cells(f'B{r}:C{r}')
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def table_sheet(wb, name, title, cols, nrows, example=None, formulas=None, calc_cols=(), widths=None, start=4):
    """cols: list of headers. formulas: {col_letter: lambda r: '=...'}; applied to example and blank rows."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    last = get_column_letter(len(cols))
    ws['A1'] = title
    ws['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    ws.row_dimensions[1].height = 26
    ws['A2'] = 'Type in white cells. Grey cells calculate automatically. See the How to use sheet.'
    ws['A2'].font = Font(name=F, size=9, color=MUTED)
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=start, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = fill(INK if get_column_letter(i) in calc_cols else GREEN)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        c.border = BORDER
    ws.row_dimensions[start].height = 34
    exs = example if (example and isinstance(example[0], list)) else ([example] if example else [])
    rows = exs + [None] * nrows
    for k, data in enumerate(rows):
        r = start + 1 + k
        for i in range(1, len(cols) + 1):
            L = get_column_letter(i)
            c = ws.cell(row=r, column=i)
            if formulas and L in formulas:
                c.value = formulas[L](r)
            elif data is not None and i - 1 < len(data):
                c.value = data[i - 1]
            c.font = Font(name=F, size=10, color=INK)
            c.border = BORDER
            c.alignment = Alignment(vertical='center')
            c.fill = EX_FILL if data is not None else (CALC_FILL if L in calc_cols else INPUT_FILL)
        ws.row_dimensions[r].height = 20
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f'A{start}:{last}{start + len(rows)}'
    ws.print_title_rows = f'{start}:{start}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = 'alhajiimehemoh.github.io'
    ws.oddFooter.right.text = 'Page &P of &N'
    return ws, start + 1, start + len(rows)


def dv_list(ws, items, rng):
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(items), allow_blank=True)
    dv.error = 'Choose a value from the list'
    dv.errorTitle = 'Not in the list'
    ws.add_data_validation(dv)
    dv.add(rng)


def status_colours(ws, rng, mapping):
    first = rng.split(':')[0]
    for text, (bg, fg) in mapping.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{text}"'], fill=fill(bg), font=Font(name=F, bold=True, color=fg)))


# ---------------------------------------------------------------------------------------------
XLSX = []


def resource(**meta):
    def wrap(fn):
        meta['build'] = fn
        XLSX.append(meta)
        return fn
    return wrap


@resource(file='stock-tracker.xlsx', short='Stock tracker', category='Inventory', title='Stock tracker',
          blurb='Opening stock, received, issued and losses give closing stock, months of stock and an automatic REORDER or OVERSTOCK flag.',
          tags=['Excel', 'Formulas'])
def _stock(path):
    wb = openpyxl.Workbook()
    cover(wb, 'Stock tracker', 'Monthly stock balance for a clinic, store or small programme, with months of stock and reorder flags.',
          ['Use one copy per store per month. Write the store name and month in the file name.',
           'List each item once in the Stock sheet with its unit, for example "tablets" or "vials".',
           'Fill opening stock, quantity received, quantity issued and any losses or adjustments.',
           'Enter the average monthly use, from the last three months if you have it.',
           'Set minimum and maximum months of stock in the Settings sheet. Status turns REORDER or OVERSTOCK by itself.'])
    st = wb.create_sheet('Settings')
    st.sheet_view.showGridLines = False
    st['A1'] = 'Settings'
    st['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    for r, (k, v, note) in enumerate([('Minimum months of stock', 2, 'Below this, Status shows REORDER'),
                                      ('Maximum months of stock', 6, 'Above this, Status shows OVERSTOCK')], 3):
        st[f'A{r}'] = k
        st[f'B{r}'] = v
        st[f'C{r}'] = note
        st[f'A{r}'].font = Font(name=F, size=10, bold=True)
        st[f'B{r}'].font = Font(name=F, size=10, color='0000FF')
        st[f'B{r}'].border = BORDER
        st[f'C{r}'].font = Font(name=F, size=9, color=MUTED)
    st['A6'] = 'These are common starting points. Use the levels your programme or national supply chain guidance sets.'
    st['A6'].font = Font(name=F, size=9, color=MUTED, italic=True)
    st.column_dimensions['A'].width = 28
    st.column_dimensions['B'].width = 10
    st.column_dimensions['C'].width = 42
    cols = ['Item', 'Unit', 'Opening stock', 'Received', 'Issued', 'Losses or adjustments', 'Closing stock',
            'Average monthly use', 'Months of stock', 'Status', 'Notes']
    fx = {'G': lambda r: f'=IF(A{r}="","",C{r}+D{r}-E{r}-F{r})',
          'I': lambda r: f'=IF(OR(A{r}="",N(H{r})=0),"",ROUND(G{r}/H{r},1))',
          'J': lambda r: f'=IF(I{r}="","",IF(I{r}<Settings!$B$3,"REORDER",IF(I{r}>Settings!$B$4,"OVERSTOCK","OK")))'}
    ws, a, b = table_sheet(wb, 'Stock', 'Monthly stock balance', cols, 40,
                           example=[['Example: Paracetamol 500 mg', 'tablets', 1200, 500, 900, 0, None, 800, None, None, 'Example row'],
                                    ['Example: ORS sachets', 'sachets', 400, 600, 300, 20, None, 250, None, None, 'Example row'],
                                    ['Example: Gloves (box of 100)', 'boxes', 90, 0, 10, 0, None, 10, None, None, 'Example row']],
                           formulas=fx, calc_cols=('G', 'I', 'J'), widths=[30, 10, 12, 11, 11, 14, 12, 14, 12, 13, 26])
    for col in 'CDEFGH':
        for r in range(a, b + 1):
            ws[f'{col}{r}'].number_format = '#,##0'
    for r in range(a, b + 1):
        ws[f'I{r}'].number_format = '0.0'
    status_colours(ws, f'J{a}:J{b}', {'REORDER': ('FDE2E1', 'B42318'), 'OVERSTOCK': (AMBER_SOFT, '7A4A00'), 'OK': (SOFT, GREEN)})
    wb.move_sheet('Settings', offset=1)
    wb.active = 1
    wb.save(path)


@resource(file='indicator-tracker.xlsx', short='Indicator tracker', category='M&E', title='Indicator tracker',
          blurb='Baseline, annual target and quarterly results for count indicators, with year-to-date, percent of target and a traffic-light status.',
          tags=['Excel', 'M&E'])
def _ind(path):
    wb = openpyxl.Workbook()
    cover(wb, 'Indicator tracker', 'Track quarterly progress against annual targets for count indicators, such as people trained or doses given.',
          ['List each indicator in the Tracker sheet with its unit and annual target.',
           'Enter each quarter\'s result as it comes in. Leave future quarters empty.',
           'Year to date, percent of target and status calculate automatically.',
           'Adjust the status thresholds in the Settings sheet to match your programme.',
           'Delete the three amber example rows when you start.',
           'Use this for counts only. For percentages and rates, report the latest value instead of adding quarters.'])
    st = wb.create_sheet('Settings')
    st.sheet_view.showGridLines = False
    st['A1'] = 'Status thresholds (share of annual target reached so far)'
    st['A1'].font = Font(name=F, size=13, bold=True, color=INK)
    for r, (k, v) in enumerate([('On track at or above', 0.9), ('Watch at or above', 0.6)], 3):
        st[f'A{r}'] = k
        st[f'B{r}'] = v
        st[f'B{r}'].number_format = '0%'
        st[f'A{r}'].font = Font(name=F, size=10, bold=True)
        st[f'B{r}'].font = Font(name=F, size=10, color='0000FF')
        st[f'B{r}'].border = BORDER
    st['A6'] = 'Expected progress: Q1 25%, Q2 50%, Q3 75%, Q4 100%. Status compares year to date with the share expected by the latest quarter entered.'
    st['A6'].font = Font(name=F, size=9, color=MUTED, italic=True)
    st.column_dimensions['A'].width = 26
    st.column_dimensions['B'].width = 10
    cols = ['Indicator', 'Unit', 'Baseline', 'Annual target', 'Q1', 'Q2', 'Q3', 'Q4', 'Year to date',
            '% of annual target', 'Expected by now', 'Status', 'Comment']
    fx = {'I': lambda r: f'=IF(A{r}="","",SUM(E{r}:H{r}))',
          'J': lambda r: f'=IF(OR(A{r}="",N(D{r})=0),"",I{r}/D{r})',
          'K': lambda r: f'=IF(COUNT(E{r}:H{r})=0,"",COUNT(E{r}:H{r})*0.25)',
          'L': lambda r: f'=IF(OR(J{r}="",K{r}=""),"",IF(J{r}/K{r}>=Settings!$B$3,"On track",IF(J{r}/K{r}>=Settings!$B$4,"Watch","Behind")))'}
    ws, a, b = table_sheet(wb, 'Tracker', 'Indicator tracker', cols, 25,
                           example=[['Example: Health workers trained in data quality', 'people', 0, 120, 30, 25, None, None, None, None, None, None, 'Example row'],
                                    ['Example: Supervision visits completed', 'visits', 0, 48, 8, 9, None, None, None, None, None, None, 'Example row'],
                                    ['Example: Facilities reporting on time', 'reports', 0, 240, 25, 20, None, None, None, None, None, None, 'Example row']],
                           formulas=fx, calc_cols=('I', 'J', 'K', 'L'), widths=[40, 10, 10, 12, 8, 8, 8, 8, 12, 12, 12, 11, 26])
    for r in range(a, b + 1):
        ws[f'J{r}'].number_format = '0%'
        ws[f'K{r}'].number_format = '0%'
        for col in 'CDEFGHI':
            ws[f'{col}{r}'].number_format = '#,##0'
    status_colours(ws, f'L{a}:L{b}', {'On track': (SOFT, GREEN), 'Watch': (AMBER_SOFT, '7A4A00'), 'Behind': ('FDE2E1', 'B42318')})
    wb.move_sheet('Settings', offset=1)
    wb.active = 1
    wb.save(path)


@resource(file='data-cleaning-log.xlsx', short='Data cleaning log', category='M&E', title='Data cleaning log',
          blurb='Record every issue found and every correction made, with drop-down issue types and an automatic summary by type and status.',
          tags=['Excel', 'Data quality'])
def _log(path):
    wb = openpyxl.Workbook()
    cover(wb, 'Data cleaning log', 'A record of every problem found in a dataset and what was done about it, so changes can always be explained.',
          ['Add one row for each problem you find, even if you fix it immediately.',
           'Use the record ID from the dataset, never a person\'s name.',
           'Choose the issue type and status from the drop-down lists.',
           'Write the original value and the corrected value, so the change can be traced.',
           'The Summary sheet counts issues by type and status automatically. Share it with your monthly report.'])
    types = ['Missing value', 'Duplicate', 'Out of range', 'Inconsistent', 'GPS error', 'Wrong date', 'Wrong code', 'Other']
    stats = ['Open', 'Query sent', 'Fixed', 'Cannot fix']
    cols = ['Date found', 'Dataset or form', 'Record ID', 'Field', 'Issue type', 'Original value', 'Corrected value',
            'Action taken', 'Status', 'Found by (role)', 'Date closed']
    ws, a, b = table_sheet(wb, 'Log', 'Data cleaning log', cols, 200,
                           example=[['2026-01-15', 'Example: EPI monthly form', 'R-0142', 'children_vaccinated', 'Out of range', '1500', '150',
                                     'Confirmed with facility: extra zero typed', 'Fixed', 'Data officer', '2026-01-16'],
                                    ['2026-01-15', 'Example: EPI monthly form', 'R-0157', 'facility', 'Duplicate', 'R-0157 and R-0158', '',
                                     'Asked facility which submission is correct', 'Query sent', 'Data officer', ''],
                                    ['2026-01-16', 'Example: Household survey', 'H-0031', 'location', 'GPS error', 'Outside district', '',
                                     'Flagged for re-visit', 'Open', 'Supervisor', '']],
                           widths=[12, 24, 12, 20, 16, 14, 14, 36, 12, 16, 12])
    dv_list(ws, types, f'E{a}:E{b}')
    dv_list(ws, stats, f'I{a}:I{b}')
    status_colours(ws, f'I{a}:I{b}', {'Open': ('FDE2E1', 'B42318'), 'Query sent': (AMBER_SOFT, '7A4A00'), 'Fixed': (SOFT, GREEN)})
    sm = wb.create_sheet('Summary')
    sm.sheet_view.showGridLines = False
    sm['A1'] = 'Summary'
    sm['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    sm['A2'] = 'Counts include the three amber example rows on the Log sheet until you delete them.'
    sm['A2'].font = Font(name=F, size=9, color=MUTED, italic=True)
    hdr = ['Issue type'] + stats + ['Total']
    for i, h in enumerate(hdr, 1):
        c = sm.cell(row=4, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = fill(GREEN)
        c.border = BORDER
    for k, t in enumerate(types):
        r = 5 + k
        sm.cell(row=r, column=1, value=t).font = Font(name=F, size=10, bold=True)
        for j, s_ in enumerate(stats):
            L = get_column_letter(2 + j)
            sm[f'{L}{r}'] = f'=COUNTIFS(Log!$E${a}:$E${b},$A{r},Log!$I${a}:$I${b},"{s_}")'
        sm[f'F{r}'] = f'=SUM(B{r}:E{r})'
        for col in 'ABCDEF':
            sm[f'{col}{r}'].border = BORDER
            sm[f'{col}{r}'].fill = CALC_FILL if col != 'A' else INPUT_FILL
            if col != 'A':
                sm[f'{col}{r}'].font = Font(name=F, size=10)
    tr = 5 + len(types)
    sm[f'A{tr}'] = 'Total'
    sm[f'A{tr}'].font = Font(name=F, size=10, bold=True)
    for col in 'BCDEF':
        sm[f'{col}{tr}'] = f'=SUM({col}5:{col}{tr - 1})'
        sm[f'{col}{tr}'].font = Font(name=F, size=10, bold=True)
        sm[f'{col}{tr}'].fill = fill(SOFT)
        sm[f'{col}{tr}'].border = BORDER
    sm.column_dimensions['A'].width = 18
    for col in 'BCDEF':
        sm.column_dimensions[col].width = 12
    wb.active = 1
    wb.save(path)


@resource(file='it-asset-register.xlsx', short='IT asset register', category='IT', title='IT asset register',
          blurb='Every computer, phone, printer and router in one list, with warranty status and next service date worked out for you.',
          tags=['Excel', 'IT'])
def _assets(path):
    wb = openpyxl.Workbook()
    cover(wb, 'IT asset register', 'One list of every device your office owns, who uses it, where it is and when it needs attention.',
          ['Give each device an asset tag, write it on a sticker on the device and enter it here.',
           'Record the user by role or office, for example "Finance officer", rather than a full name if you prefer.',
           'Enter purchase and warranty end dates as dates. Warranty status updates every day by itself.',
           'After each service, update the Last service date. The next service date is set 6 months later.',
           'Review the list every quarter and mark lost, retired or disposed devices instead of deleting them.'])
    cols = ['Asset tag', 'Type', 'Make and model', 'Serial number', 'Assigned to (role)', 'Location', 'Purchase date',
            'Warranty end', 'Warranty status', 'Last service', 'Next service', 'Condition', 'Notes']
    fx = {'I': lambda r: f'=IF(H{r}="","",IF(H{r}<TODAY(),"Expired",IF(H{r}-TODAY()<=60,"Ending soon","Active")))',
          'K': lambda r: f'=IF(J{r}="","",EDATE(J{r},6))'}
    import datetime as dt
    ws, a, b = table_sheet(wb, 'Assets', 'IT asset register', cols, 100,
                           example=[['EX-001', 'Laptop', 'Example: brand and model', 'SN-EXAMPLE1', 'Data officer', 'Main office',
                                     dt.date(2025, 3, 1), dt.date(2028, 3, 1), None, dt.date(2026, 6, 1), None, 'Good', 'Example row'],
                                    ['EX-002', 'Printer', 'Example: brand and model', 'SN-EXAMPLE2', 'Shared', 'Main office',
                                     dt.date(2022, 5, 10), dt.date(2024, 5, 10), None, dt.date(2026, 2, 1), None, 'Needs repair', 'Example row'],
                                    ['EX-003', 'Router', 'Example: brand and model', 'SN-EXAMPLE3', 'Shared', 'Server corner',
                                     dt.date(2024, 11, 20), None, None, dt.date(2026, 9, 15), None, 'Good', 'Example row']],
                           formulas=fx, calc_cols=('I', 'K'), widths=[11, 13, 24, 16, 18, 14, 13, 13, 14, 13, 13, 12, 24])
    for r in range(a, b + 1):
        for col in 'GHJK':
            ws[f'{col}{r}'].number_format = 'dd mmm yyyy'
    dv_list(ws, ['Desktop', 'Laptop', 'Tablet', 'Phone', 'Printer', 'Scanner', 'Router', 'Switch', 'UPS', 'Projector', 'Other'], f'B{a}:B{b}')
    dv_list(ws, ['Good', 'Needs repair', 'Retired', 'Lost', 'Disposed'], f'L{a}:L{b}')
    status_colours(ws, f'I{a}:I{b}', {'Expired': ('FDE2E1', 'B42318'), 'Ending soon': (AMBER_SOFT, '7A4A00'), 'Active': (SOFT, GREEN)})
    wb.active = 1
    wb.save(path)


@resource(file='sample-epi-xlsform.xlsx', short='Sample XLSForm', category='Data collection', title='Sample EPI XLSForm',
          blurb='A short monthly facility form with validation, skip logic and a calculated field. Upload to Survey123 Connect or KoboToolbox and adapt.',
          tags=['XLSForm', 'Survey123', 'Kobo'])
def _form(path):
    """A real XLSForm: the survey, choices and settings sheets must keep their exact names and columns."""
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'survey'
    s.append(['type', 'name', 'label', 'hint', 'required', 'relevant', 'constraint', 'constraint_message', 'calculation', 'appearance'])
    rows = [
        ['start', 'start'], ['end', 'end'],
        ['note', 'intro', 'Monthly EPI facility check (sample form). Do not record names or phone numbers.'],
        ['select_one district', 'district', 'District', '', 'yes', '', '', '', '', 'minimal'],
        ['select_one facility', 'facility', 'Health facility', 'Only facilities in the chosen district are listed', 'yes', '', '', '', '', 'minimal'],
        ['date', 'report_month', 'Reporting month', 'Pick any day in the month', 'yes', '', '. <= today()', 'The month cannot be in the future', '', 'month-year'],
        ['begin group', 'doses', 'Doses given this month', '', '', '', '', '', '', 'field-list'],
        ['integer', 'penta1', 'Penta 1 doses given', '', 'yes', '', '. >= 0 and . <= 2000', 'Enter a number from 0 to 2000'],
        ['integer', 'penta3', 'Penta 3 doses given', '', 'yes', '', '. >= 0 and . <= 2000', 'Enter a number from 0 to 2000'],
        ['end group', 'doses'],
        ['calculate', 'dropout', '', '', '', '', '', '', 'if(${penta1} > 0, round((${penta1} - ${penta3}) div ${penta1} * 100, 1), 0)'],
        ['note', 'dropout_note', 'Penta 1 to Penta 3 dropout: ${dropout}%'],
        ['select_one yesno', 'stock_out', 'Any vaccine stock-out this month?', '', 'yes'],
        ['select_multiple vaccine', 'stock_out_which', 'Which vaccines ran out?', 'Tick all that apply', 'yes', "${stock_out} = 'yes'"],
        ['integer', 'stock_out_days', 'Total days without stock', '', 'yes', "${stock_out} = 'yes'", '. >= 1 and . <= 31', 'Enter 1 to 31 days'],
        ['geopoint', 'location', 'Facility location', 'Capture at the facility', 'no'],
        ['text', 'comments', 'Comments', '', 'no', '', '', '', '', 'multiline'],
    ]
    for r in rows:
        s.append(r)
    c = wb.create_sheet('choices')
    c.append(['list_name', 'name', 'label', 'district'])
    for r in [['district', 'dist_a', 'Sample District A'], ['district', 'dist_b', 'Sample District B'],
              ['facility', 'phc_a1', 'Sample PHC A1', 'dist_a'], ['facility', 'phc_a2', 'Sample PHC A2', 'dist_a'],
              ['facility', 'phc_b1', 'Sample PHC B1', 'dist_b'], ['facility', 'phc_b2', 'Sample PHC B2', 'dist_b'],
              ['yesno', 'yes', 'Yes'], ['yesno', 'no', 'No'],
              ['vaccine', 'bcg', 'BCG'], ['vaccine', 'opv', 'OPV'], ['vaccine', 'penta', 'Pentavalent'], ['vaccine', 'pcv', 'PCV'],
              ['vaccine', 'rota', 'Rotavirus'], ['vaccine', 'mr', 'Measles-Rubella'], ['vaccine', 'other', 'Other']]:
        c.append(r)
    for i, row in enumerate(s.iter_rows(min_row=1, max_row=s.max_row), 1):
        if row[0].value == 'select_one facility':
            s.cell(row=i, column=11, value='district=${district}')
    s.cell(row=1, column=11, value='choice_filter')
    st = wb.create_sheet('settings')
    st.append(['form_title', 'form_id', 'version'])
    st.append(['Sample EPI monthly facility check', 'sample_epi_check', '1'])
    for sh in (s, c, st):
        for cell in sh[1]:
            cell.font = Font(name=F, bold=True, color='FFFFFF')
            cell.fill = fill(GREEN)
        sh.freeze_panes = 'A2'
        for col in range(1, sh.max_column + 1):
            sh.column_dimensions[get_column_letter(col)].width = 24
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name=F, size=10)
    s.column_dimensions['C'].width = 44
    s.column_dimensions['I'].width = 60
    for i, row in enumerate(s.iter_rows(min_row=2), 2):
        t = str(row[0].value or '')
        if t.startswith(('begin', 'end group')):
            for cell in row:
                cell.fill = fill(SOFT)
    wb.save(path)


@resource(file='reporting-timeliness-tracker.xlsx', short='Reporting tracker', category='M&E', title='Reporting timeliness tracker',
          blurb='Log the date each facility report arrives. Every month turns On time, Late or Missing, with completeness and timeliness rates and a chart.',
          tags=['Excel', 'Chart', 'M&E'])
def _timely(path):
    import datetime as dt
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.Workbook()
    cover(wb, 'Reporting timeliness tracker', 'See at a glance which facilities report on time, which are late and which have not reported, month by month.',
          ['Set the year and the monthly deadline day in the Settings sheet.',
           'List each facility once in column A of the Received sheet.',
           'When a report arrives, type the date it arrived under that month.',
           'The Status sheet shows On time, Late or Missing. Missing appears only after the deadline has passed.',
           'The Summary sheet gives completeness and timeliness rates per month and a chart. Delete the three amber example facilities when you start.'])
    st = wb.create_sheet('Settings')
    st.sheet_view.showGridLines = False
    st['A1'] = 'Settings'
    st['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    for r, (k, v, note) in enumerate([('Reporting year', 2026, 'The year the reports cover'),
                                      ('Deadline day of the following month', 5, 'Example: 5 means the January report is due by 5 February')], 3):
        st[f'A{r}'] = k
        st[f'B{r}'] = v
        st[f'C{r}'] = note
        st[f'A{r}'].font = Font(name=F, size=10, bold=True)
        st[f'B{r}'].font = Font(name=F, size=10, color='0000FF')
        st[f'B{r}'].border = BORDER
        st[f'C{r}'].font = Font(name=F, size=9, color=MUTED)
    st.column_dimensions['A'].width = 36
    st.column_dimensions['B'].width = 10
    st.column_dimensions['C'].width = 56
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    # example arrival dates: day of the following month, None = not received
    ex = {'Example: Facility A': [3, 4, 2, 5, 3, 4, 2, 3, 4], 'Example: Facility B': [6, 4, 9, 5, 12, 4, 5, 7, 5],
          'Example: Facility C': [4, None, 5, 3, None, 5, 6, 4, None]}
    rows = []
    for name, days in ex.items():
        row = [name]
        for m, d in enumerate(days, 1):
            row.append(None if d is None else dt.date(2026 + (m == 12), m % 12 + 1, d))
        rows.append(row + [None] * (12 - len(days)))
    ws, a, b = table_sheet(wb, 'Received', 'Date each report was received', ['Facility'] + months, 27, example=rows,
                           widths=[28] + [11] * 12)
    for r in range(a, b + 1):
        for c in range(2, 14):
            ws.cell(row=r, column=c).number_format = 'dd mmm'
    fx = {}
    for i, m in enumerate(months, 2):
        L = get_column_letter(i)
        due = f'DATE(Settings!$B$3,{i},Settings!$B$4)'
        fx[L] = (lambda L, due: (lambda r: f'=IF(Received!$A{r}="","",IF(Received!{L}{r}="",IF(TODAY()>{due},"Missing",""),IF(Received!{L}{r}<={due},"On time","Late")))'))(L, due)
    fx['A'] = lambda r: f'=IF(Received!A{r}="","",Received!A{r})'
    ws2, a2, b2 = table_sheet(wb, 'Status', 'Status by month', ['Facility'] + months, 30, formulas=fx,
                              calc_cols=tuple(get_column_letter(i) for i in range(1, 14)), widths=[28] + [10] * 12)
    for r in range(a2, a2 + 3):
        for c in range(1, 14):
            ws2.cell(row=r, column=c).fill = CALC_FILL
    status_colours(ws2, f'B{a2}:M{b2}', {'Missing': ('FDE2E1', 'B42318'), 'Late': (AMBER_SOFT, '7A4A00'), 'On time': (SOFT, GREEN)})
    sm = wb.create_sheet('Summary')
    sm.sheet_view.showGridLines = False
    sm['A1'] = 'Completeness and timeliness'
    sm['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    sm['A2'] = 'Completeness: reports received out of facilities listed. Timeliness: reports on time out of facilities listed. Months not yet due stay blank.'
    sm['A2'].font = Font(name=F, size=9, color=MUTED, italic=True)
    hdr = ['Month', 'Facilities', 'Received', 'On time', 'Late', 'Missing', 'Completeness', 'Timeliness']
    for i, h in enumerate(hdr, 1):
        c = sm.cell(row=4, column=i, value=h)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = fill(GREEN)
        c.border = BORDER
    for k, m in enumerate(months):
        r = 5 + k
        L = get_column_letter(2 + k)
        due = f'DATE(Settings!$B$3,{k + 2},Settings!$B$4)'
        sm[f'A{r}'] = m
        sm[f'B{r}'] = f'=COUNTA(Received!$A${a}:$A${b})'
        sm[f'C{r}'] = f'=COUNT(Received!{L}${a}:{L}${b})'
        sm[f'D{r}'] = f'=COUNTIF(Status!{L}${a2}:{L}${b2},"On time")'
        sm[f'E{r}'] = f'=COUNTIF(Status!{L}${a2}:{L}${b2},"Late")'
        sm[f'F{r}'] = f'=COUNTIF(Status!{L}${a2}:{L}${b2},"Missing")'
        sm[f'G{r}'] = f'=IF(OR(B{r}=0,AND(TODAY()<={due},C{r}=0)),"",C{r}/B{r})'
        sm[f'H{r}'] = f'=IF(G{r}="","",D{r}/B{r})'
        for col in 'ABCDEFGH':
            cell = sm[f'{col}{r}']
            cell.border = BORDER
            cell.font = Font(name=F, size=10, bold=(col == 'A'))
            cell.fill = INPUT_FILL if col == 'A' else CALC_FILL
        sm[f'G{r}'].number_format = '0%'
        sm[f'H{r}'].number_format = '0%'
    for col, w in zip('ABCDEFGH', [10, 11, 11, 10, 8, 10, 14, 12]):
        sm.column_dimensions[col].width = w
    ch = BarChart()
    ch.type = 'col'
    ch.title = 'Completeness and timeliness by month'
    ch.y_axis.title = 'Share of facilities'
    ch.y_axis.number_format = '0%'
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.add_data(Reference(sm, min_col=7, max_col=8, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(sm, min_col=1, min_row=5, max_row=16))
    ch.height = 8
    ch.width = 18
    s1, s2 = ch.series
    s1.graphicalProperties.solidFill = GREEN
    s2.graphicalProperties.solidFill = AMBER
    sm.add_chart(ch, 'J4')
    wb.move_sheet('Settings', offset=3)
    wb.active = 1
    wb.save(path)


@resource(file='training-tracker.xlsx', short='Training tracker', category='Training', title='Training attendance and results tracker',
          blurb='Attendance for up to five days, pre and post test scores, improvement and certificate eligibility, with a summary and a chart.',
          tags=['Excel', 'Chart', 'Training'])
def _training(path):
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.Workbook()
    cover(wb, 'Training attendance and results tracker', 'Track who attended each day of a training, how much they learned and who qualifies for a certificate.',
          ['Fill the training name, dates and rules in the Settings sheet.',
           'Give each participant an ID. Names are optional; use IDs if the file will be shared.',
           'Mark attendance for each day as Y or N from the drop-down.',
           'Enter pre-test and post-test scores as percentages, for example 45%.',
           'Days attended, improvement and certificate eligibility calculate automatically. Delete the three amber example rows when you start.'])
    st = wb.create_sheet('Settings')
    st.sheet_view.showGridLines = False
    st['A1'] = 'Settings'
    st['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    for r, (k, v, fmt) in enumerate([('Training name', 'Example: Data quality for facility staff', None),
                                     ('Number of training days (1 to 5)', 3, None),
                                     ('Minimum days for a certificate', 3, None),
                                     ('Minimum post-test score for a certificate', 0.7, '0%')], 3):
        st[f'A{r}'] = k
        st[f'B{r}'] = v
        st[f'A{r}'].font = Font(name=F, size=10, bold=True)
        st[f'B{r}'].font = Font(name=F, size=10, color='0000FF')
        st[f'B{r}'].border = BORDER
        if fmt:
            st[f'B{r}'].number_format = fmt
    st.column_dimensions['A'].width = 40
    st.column_dimensions['B'].width = 40
    cols = ['Participant ID', 'Name (optional)', 'Organisation or facility', 'Role', 'Sex', 'Day 1', 'Day 2', 'Day 3', 'Day 4', 'Day 5',
            'Days attended', 'Pre-test', 'Post-test', 'Improvement (points)', 'Certificate']
    fx = {'K': lambda r: f'=IF(A{r}="","",COUNTIF(F{r}:J{r},"Y"))',
          'N': lambda r: f'=IF(OR(A{r}="",L{r}="",M{r}=""),"",(M{r}-L{r})*100)',
          'O': lambda r: f'=IF(A{r}="","",IF(AND(K{r}>=Settings!$B$5,N(M{r})>=Settings!$B$6),"Yes","No"))'}
    ws, a, b = table_sheet(wb, 'Participants', 'Participants', cols, 60,
                           example=[['EX-01', '', 'Example: PHC A', 'Nurse', 'F', 'Y', 'Y', 'Y', '', '', None, 0.45, 0.8, None, None],
                                    ['EX-02', '', 'Example: PHC B', 'Data clerk', 'M', 'Y', 'N', 'Y', '', '', None, 0.5, 0.75, None, None],
                                    ['EX-03', '', 'Example: PHC C', 'CHW supervisor', 'F', 'Y', 'Y', 'Y', '', '', None, 0.35, 0.65, None, None]],
                           formulas=fx, calc_cols=('K', 'N', 'O'), widths=[13, 18, 22, 16, 7, 7, 7, 7, 7, 7, 10, 10, 10, 13, 12])
    for r in range(a, b + 1):
        ws[f'L{r}'].number_format = '0%'
        ws[f'M{r}'].number_format = '0%'
        ws[f'N{r}'].number_format = '+0;-0;0'
    dv_list(ws, ['Y', 'N'], f'F{a}:J{b}')
    dv_list(ws, ['F', 'M'], f'E{a}:E{b}')
    status_colours(ws, f'O{a}:O{b}', {'Yes': (SOFT, GREEN), 'No': ('FDE2E1', 'B42318')})
    sm = wb.create_sheet('Summary')
    sm.sheet_view.showGridLines = False
    sm['A1'] = 'Training summary'
    sm['A1'].font = Font(name=F, size=16, bold=True, color=INK)
    sm['A2'] = '=Settings!B3'
    sm['A2'].font = Font(name=F, size=10, color=MUTED)
    R = f'${a}:$'
    rows = [('Participants', f'=COUNTA(Participants!$A${a}:$A${b})', '0'),
            ('Female', f'=COUNTIF(Participants!$E${a}:$E${b},"F")', '0'),
            ('Male', f'=COUNTIF(Participants!$E${a}:$E${b},"M")', '0'),
            ('Average days attended', f'=IFERROR(AVERAGE(Participants!$K${a}:$K${b}),"")', '0.0'),
            ('Average pre-test', f'=IFERROR(AVERAGE(Participants!$L${a}:$L${b}),"")', '0%'),
            ('Average post-test', f'=IFERROR(AVERAGE(Participants!$M${a}:$M${b}),"")', '0%'),
            ('Average improvement (points)', f'=IFERROR(AVERAGE(Participants!$N${a}:$N${b}),"")', '+0.0;-0.0;0'),
            ('Eligible for a certificate', f'=COUNTIF(Participants!$O${a}:$O${b},"Yes")', '0')]
    for k, (lab, f_, fmt) in enumerate(rows, 4):
        sm[f'A{k}'] = lab
        sm[f'B{k}'] = f_
        sm[f'A{k}'].font = Font(name=F, size=10, bold=True)
        sm[f'B{k}'].font = Font(name=F, size=12, bold=True, color=GREEN)
        sm[f'B{k}'].number_format = fmt
        sm[f'B{k}'].fill = CALC_FILL
        for col in 'AB':
            sm[f'{col}{k}'].border = BORDER
    sm.column_dimensions['A'].width = 32
    sm.column_dimensions['B'].width = 14
    ch = BarChart()
    ch.type = 'col'
    ch.title = 'Average score before and after'
    ch.y_axis.number_format = '0%'
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.add_data(Reference(sm, min_col=2, min_row=8, max_row=9))
    ch.set_categories(Reference(sm, min_col=1, min_row=8, max_row=9))
    ch.legend = None
    ch.height = 7
    ch.width = 12
    ch.series[0].graphicalProperties.solidFill = GREEN
    sm.add_chart(ch, 'D4')
    wb.move_sheet('Settings', offset=2)
    wb.active = 1
    wb.save(path)


def build_xlsx(outdir):
    os.makedirs(outdir, exist_ok=True)
    for r in XLSX:
        r['build'](os.path.join(outdir, r['file']))
    return XLSX
